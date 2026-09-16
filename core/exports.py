# core/exports.py
"""Newsletter intermediate structure plus PDF/WORD/print/share writers."""
from __future__ import annotations

import logging
import os
import subprocess
import platform
import webbrowser
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable
from xml.sax.saxutils import escape

from core.briefs import build_brief, hard_trim
from core.scoring import area_of, relevance

logger = logging.getLogger(__name__)

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
except ImportError:  # pragma: no cover - import error surfaces in Task 1 install
    A4 = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None


def _section_body(doc: dict, api_client) -> str:
    """Brief body (single paragraph, no Fonte line); never raises."""
    try:
        body, _, _ = build_brief(doc, api_client).rpartition("\n")
        return body
    except Exception as exc:  # last-resort guard: export never fails
        logger.warning("Brief build failed, using trimmed summary: %s", exc)
        return hard_trim(str(doc.get("summary") or doc.get("overview") or ""))


def build_newsletter(docs: list[dict], weights: dict, mode: str, meta: dict,
                      preserve_order: bool = False, *,
                      api_client=None,
                      progress_cb: Callable[[int, int], None] | None = None,
                      include_key_points: bool = False) -> dict:
    """Rank ``docs`` by relevance into newsletter sections.

    With ``preserve_order=True`` the input order is kept as-is (used by
    the Personalizado flow, whose modal order is authoritative).

    Each section carries the brief (single pt-BR paragraph when an
    ``api_client`` is available, else an extractive fallback) in
    ``summary``; the ``Fonte: {url}`` line is rendered from ``url`` by
    the writers. ``progress_cb(i, total)`` is called per document (used
    for the "Condensando resumo i/N..." status). ``include_key_points``
    re-adds the ``key_points`` bullets AFTER the Fonte line.
    """
    ranked = list(docs) if preserve_order else sorted(
        docs, key=lambda d: relevance(d, weights), reverse=True)
    total = len(ranked)
    sections = []
    for i, d in enumerate(ranked, start=1):
        if progress_cb is not None:
            try:
                progress_cb(i, total)
            except Exception:
                logger.debug("Newsletter progress callback failed",
                             exc_info=True)
        sections.append({
            "title": d.get("newsletter") or d.get("title", ""),
            "summary": _section_body(d, api_client),
            "key_points": list(d.get("key_points") or []),
            "area": area_of(d),
            "relevance": relevance(d, weights),
            "organizations": [o.get("name", "") for o in d.get("organization") or []],
            "countries": list(d.get("related_country") or []),
            "url": d.get("url", ""),
        })
    return {"title": meta.get("title", "GLOBAL QUANTUM INTELLIGENCE – QuIIN"),
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "mode": mode, "weights": dict(weights),
            "include_key_points": bool(include_key_points),
            "sections": sections}


def _with_key_points(structure: dict) -> bool:
    return bool(structure.get("include_key_points"))


def export_pdf(structure: dict, path: Path | str) -> str:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(structure["title"]), styles["Title"]), Spacer(1, 12)]
    show_points = _with_key_points(structure)
    for section in structure["sections"]:
        story.append(Paragraph(f"{escape(section['title'])} — {escape(section['area'])}"
                               f" ({section['relevance']})", styles["Heading2"]))
        story.append(Paragraph(escape(section["summary"]), styles["BodyText"]))
        story.append(Paragraph(f"Fonte: {escape(section['url'])}",
                               styles["BodyText"]))
        if show_points:
            for point in section["key_points"]:
                story.append(Paragraph(f"• {escape(point)}", styles["BodyText"]))
        story.append(Spacer(1, 12))
    SimpleDocTemplate(str(target), pagesize=A4).build(story)
    return str(target)


def export_word(structure: dict, path: Path | str) -> str:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(structure["title"], level=0)
    show_points = _with_key_points(structure)
    for section in structure["sections"]:
        document.add_heading(f"{section['title']} — {section['area']}"
                             f" ({section['relevance']})", level=1)
        document.add_paragraph(section["summary"])
        document.add_paragraph(f"Fonte: {section['url']}")
        if show_points:
            for point in section["key_points"]:
                document.add_paragraph(point, style="List Bullet")
    document.save(str(target))
    return str(target)


def print_pdf(path: Path | str) -> str:
    if platform.system() == "Windows":
        try:
            os.startfile(str(path), "print")  # noqa: S606
            return "printed"
        except OSError:
            pass
    webbrowser.open(str(path))
    return "opened"


def share_package(structure: dict, dir: Path | str) -> str:
    target_dir = Path(dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    lines = [f"# {structure['title']}", ""]
    show_points = _with_key_points(structure)
    for section in structure["sections"]:
        lines.append(f"## {section['title']} — {section['area']} ({section['relevance']})")
        lines.append("")
        lines.append(section["summary"])
        lines.append("")
        lines.append(f"Fonte: {section['url']}")
        lines.append("")
        if show_points:
            lines.extend(f"- {point}" for point in section["key_points"])
            lines.append("")
    md_path = target_dir / "newsletter.md"
    md_path.write_text("\n".join(lines), encoding="utf-8")
    try:
        if platform.system() == "Windows":
            os.startfile(str(target_dir))  # noqa: S606
        elif platform.system() == "Darwin":
            subprocess.Popen(["open", str(target_dir)])
        else:
            subprocess.Popen(["xdg-open", str(target_dir)])
    except (OSError, subprocess.SubprocessError):
        pass
    return str(md_path)
